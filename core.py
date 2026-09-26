"""业务逻辑：抽取、撤回、重置、名单管理、导入导出。

所有函数都接收一个 sqlite3.Connection，方便测试时用临时库。
抽取类操作用 `with conn:` 包成单个事务，避免出现
「学生已移出池子但记录没落库」这类中间状态。
"""
from __future__ import annotations

import csv
import io
import json
import re
import secrets
import sqlite3
from datetime import datetime
from pathlib import Path

import db

# 单列粘贴时的行首序号，如「1. 张三」「2、李四」。
# 必须带分隔符才剥离，否则会把「3班李四」这类名字削掉首字。
_NUMBERING = re.compile(r"^\s*\d+\s*[.、,，)）:：]\s*")

CSV_HEADER = ["序号", "日期", "时间", "学期", "姓名", "学号", "班级"]

HEADER_ALIASES = {
    "name": {"姓名", "名字", "学生", "学生姓名", "name"},
    "sid": {"学号", "编号", "学籍号", "考号", "sid", "id"},
    "cls": {"班级", "班", "行政班", "组别", "分组", "组", "class"},
}


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def split_stamp(stamp: str) -> tuple[str, str]:
    """把「2026-09-17 08:03:12」拆成（日期, 时间）。

    日期单独成列，方便按天统计「每天抽了谁」。格式不符合预期时
    整体放进日期列，避免丢数据。
    """
    s = str(stamp or "")
    if len(s) >= 19 and s[10] == " ":
        return s[:10], s[11:]
    return s, ""


# ---------------------------------------------------------------- meta

def get_round(conn: sqlite3.Connection) -> int:
    row = conn.execute("SELECT value FROM meta WHERE key='round'").fetchone()
    return int(row["value"]) if row else 1


def set_round(conn: sqlite3.Connection, value: int) -> None:
    conn.execute(
        "INSERT INTO meta (key, value) VALUES ('round', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (str(value),),
    )


# ---------------------------------------------------------------- 查询

def stats(conn: sqlite3.Connection) -> dict:
    """界面顶部要显示的计数。"""
    total = conn.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    enabled = conn.execute(
        "SELECT COUNT(*) FROM students WHERE enabled=1"
    ).fetchone()[0]
    remaining = conn.execute(
        "SELECT COUNT(*) FROM students WHERE enabled=1 AND drawn=0"
    ).fetchone()[0]
    disabled = total - enabled
    return {
        "total": total,
        "enabled": enabled,
        "remaining": remaining,
        "drawn": enabled - remaining,
        "disabled": disabled,
        "round": get_round(conn),
    }


def list_students(conn: sqlite3.Connection) -> list[dict]:
    rows = conn.execute(
        "SELECT id, name, sid, cls, enabled, drawn, drawn_at "
        "FROM students ORDER BY cls, sid, id"
    ).fetchall()
    return [dict(r) for r in rows]


def pool_names(conn: sqlite3.Connection) -> list[str]:
    """仅取池中姓名，供界面做滚动动画使用。"""
    rows = conn.execute(
        "SELECT name FROM students WHERE enabled=1 AND drawn=0"
    ).fetchall()
    return [r["name"] for r in rows]


def history(conn: sqlite3.Connection, round: int | None = None) -> list[dict]:
    if round is None:
        rows = conn.execute(
            "SELECT * FROM draws ORDER BY round DESC, seq ASC"
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM draws WHERE round=? ORDER BY seq ASC", (round,)
        ).fetchall()
    return [dict(r) for r in rows]


def all_rounds(conn: sqlite3.Connection) -> list[int]:
    rows = conn.execute(
        "SELECT DISTINCT round FROM draws ORDER BY round DESC"
    ).fetchall()
    return [r["round"] for r in rows]


# ---------------------------------------------------------------- 抽取

def draw(conn: sqlite3.Connection, now=None, min_minutes=None,
         ignore_schedule: bool = False) -> dict | None:
    """随机抽一名学生并立即将其移出抽取池。

    池子 = 未停用 AND 本学期未谈过 AND 此刻没有课。
    now 可注入，便于测试固定时刻；未录课表时退化为不排除任何人。
    返回被抽中的学生信息；池子为空时返回 None。
    """
    with conn:
        pool = pool_status(conn, now=now, min_minutes=min_minutes,
                           ignore_schedule=ignore_schedule)["available"]
        if not pool:
            return None

        picked = secrets.choice(pool)
        rnd = get_round(conn)
        seq = conn.execute(
            "SELECT COALESCE(MAX(seq), 0) + 1 FROM draws WHERE round=?", (rnd,)
        ).fetchone()[0]
        stamp = now_str()

        cur = conn.execute(
            "INSERT INTO draws (student_id, seq, round, drawn_at, name, sid, cls) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (picked["id"], seq, rnd, stamp, picked["name"], picked["sid"], picked["cls"]),
        )
        conn.execute(
            "UPDATE students SET drawn=1, drawn_at=? WHERE id=?",
            (stamp, picked["id"]),
        )

        return {
            "draw_id": cur.lastrowid,
            "student_id": picked["id"],
            "seq": seq,
            "round": rnd,
            "name": picked["name"],
            "sid": picked["sid"],
            "cls": picked["cls"],
            "drawn_at": stamp,
            "remaining": conn.execute(
                "SELECT COUNT(*) FROM students WHERE enabled=1 AND drawn=0"
            ).fetchone()[0],
        }


def _renumber_round(conn: sqlite3.Connection, round_no: int) -> None:
    """撤回后重排本轮序号，避免出现 1,2,4,5 这样的空洞。"""
    rows = conn.execute(
        "SELECT id FROM draws WHERE round=? ORDER BY drawn_at, id", (round_no,)
    ).fetchall()
    for i, row in enumerate(rows, 1):
        conn.execute("UPDATE draws SET seq=? WHERE id=?", (i, row["id"]))


def undo(conn: sqlite3.Connection, draw_id: int) -> dict:
    """撤回一次抽取：删除记录并把该学生放回池子。"""
    with conn:
        row = conn.execute("SELECT * FROM draws WHERE id=?", (draw_id,)).fetchone()
        if row is None:
            raise ValueError("该抽取记录已不存在")

        conn.execute("DELETE FROM draws WHERE id=?", (draw_id,))

        restored = False
        if row["student_id"] is not None:
            cur = conn.execute(
                "UPDATE students SET drawn=0, drawn_at=NULL "
                "WHERE id=? AND enabled=1",
                (row["student_id"],),
            )
            restored = cur.rowcount > 0

        _renumber_round(conn, row["round"])
        remaining = conn.execute(
            "SELECT COUNT(*) FROM students WHERE enabled=1 AND drawn=0"
        ).fetchone()[0]

        return {
            "name": row["name"],
            "restored": restored,
            "remaining": remaining,
        }


def start_new_term(conn: sqlite3.Connection) -> dict:
    """开始新学期：所有人回到抽取池，历史归档保留，学期 +1。

    复用原来的 round 字段当学期计数（语义一样，都是「一个可重置的批次」），
    因此不需要改数据库结构。

    只清除「已抽出」标记，不动「手动停用」的学生，
    否则请假/转走的人会被错误地放回池子。
    """
    with conn:
        prev = get_round(conn)
        cur = conn.execute(
            "UPDATE students SET drawn=0, drawn_at=NULL WHERE drawn=1"
        )
        new_round = prev + 1
        set_round(conn, new_round)
        return {"prev_round": prev, "new_round": new_round,
                "prev_term": prev, "new_term": new_round,
                "restored": cur.rowcount}


# ---------------------------------------------------------------- 作息与课表

def _parse_hm(value) -> tuple[int, int]:
    """"08:00" -> (8, 0)。数据坏了就退回 (0, 0)，不让判定崩掉。"""
    try:
        parts = str(value).split(":")
        return int(parts[0]), int(parts[1])
    except Exception:
        return 0, 0


def _minute_of_day(value) -> int:
    h, m = _parse_hm(value)
    return h * 60 + m


def get_min_talk_minutes(conn: sqlite3.Connection) -> int:
    """离下个有课的大课不足这个分钟数，就不算空闲。"""
    row = conn.execute(
        "SELECT value FROM meta WHERE key='min_talk_minutes'"
    ).fetchone()
    try:
        return max(0, int(row["value"])) if row else db.DEFAULT_MIN_TALK_MINUTES
    except (TypeError, ValueError):
        return db.DEFAULT_MIN_TALK_MINUTES


def set_min_talk_minutes(conn: sqlite3.Connection, minutes) -> dict:
    value = max(0, int(minutes))
    with conn:
        conn.execute(
            "INSERT INTO meta (key, value) VALUES ('min_talk_minutes', ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (str(value),),
        )
    return {"min_talk_minutes": value}


def list_blocks(conn: sqlite3.Connection) -> list[dict]:
    """大课列表，按开始时间排序（不按 idx，避免用户改乱时间后判定出错）。"""
    rows = conn.execute(
        "SELECT idx, name, start, end FROM blocks"
    ).fetchall()
    out = [dict(r) for r in rows]
    out.sort(key=lambda b: (_minute_of_day(b["start"]), b["idx"]))
    return out


def set_block(conn: sqlite3.Connection, idx, name=None, start=None, end=None) -> dict:
    fields, values = [], []
    for col, val in (("name", name), ("start", start), ("end", end)):
        if val is not None:
            fields.append(f"{col}=?")
            values.append(str(val).strip())
    if not fields:
        return {"updated": 0}
    with conn:
        values.append(int(idx))
        cur = conn.execute(
            f"UPDATE blocks SET {', '.join(fields)} WHERE idx=?", values
        )
        return {"updated": cur.rowcount}


def schedule_configured(conn: sqlite3.Connection) -> bool:
    """是否已经录过课表。

    没录课表时必须退化成「不排除任何人」——否则用户还没录课表
    就会发现一个人也抽不出来。这是兼容性底线。
    """
    if conn.execute(
        "SELECT 1 FROM schedule WHERE has_class=1 LIMIT 1"
    ).fetchone():
        return True
    if conn.execute(
        "SELECT 1 FROM schedule_exceptions WHERE has_class=1 LIMIT 1"
    ).fetchone():
        return True
    return False


def schedule_grid(conn: sqlite3.Connection) -> dict:
    """课表网格：行是班级，列是「星期 × 大课」。"""
    classes = [r["cls"] for r in conn.execute(
        "SELECT DISTINCT cls FROM students WHERE cls <> '' ORDER BY cls"
    )]
    # 课表里有、但名单里已经没了的班级也列出来，避免脏数据看不见
    for r in conn.execute("SELECT DISTINCT cls FROM schedule WHERE cls <> ''"):
        if r["cls"] not in classes:
            classes.append(r["cls"])

    cells = {}
    for r in conn.execute(
        "SELECT cls, weekday, block, has_class, weeks, note FROM schedule"
    ):
        cells[f"{r['cls']}|{r['weekday']}|{r['block']}"] = {
            "has_class": r["has_class"], "weeks": r["weeks"] or "", "note": r["note"],
        }
    return {
        "classes": classes,
        "blocks": list_blocks(conn),
        "cells": cells,
        "configured": schedule_configured(conn),
        "termStart": get_term_start(conn),
    }


def set_schedule(conn: sqlite3.Connection, cls, weekday, block,
                 has_class=True, note="", weeks="") -> dict:
    """设置某班某天某个大课有没有课。空着且无备注就删掉该行，保持数据干净。"""
    cls = str(cls or "").strip()
    if not cls:
        raise ValueError("班级不能为空")
    weekday, block = int(weekday), int(block)
    note = str(note or "").strip()
    weeks = normalize_weeks(weeks)
    with conn:
        if not has_class and not note:
            cur = conn.execute(
                "DELETE FROM schedule WHERE cls=? AND weekday=? AND block=?",
                (cls, weekday, block),
            )
            return {"deleted": cur.rowcount}
        conn.execute(
            "INSERT INTO schedule (cls, weekday, block, has_class, weeks, note) "
            "VALUES (?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(cls, weekday, block) DO UPDATE SET "
            "has_class=excluded.has_class, weeks=excluded.weeks, note=excluded.note",
            (cls, weekday, block, 1 if has_class else 0, weeks, note),
        )
    return {"ok": True}


def list_exceptions(conn: sqlite3.Connection) -> list[dict]:
    rows = conn.execute(
        "SELECT e.student_id, s.name, s.cls, e.weekday, e.block, "
        "       e.has_class, e.note "
        "FROM schedule_exceptions e LEFT JOIN students s ON s.id = e.student_id "
        "ORDER BY s.name, e.weekday, e.block"
    ).fetchall()
    return [dict(r) for r in rows]


def set_exception(conn: sqlite3.Connection, student_id, weekday, block,
                  has_class=True, note="") -> dict:
    """个别学生的课表例外：额外有课(1) 或 全班有课但他不上(0)。

    注意这里与班级课表不同：has_class=0 表示「全班有课，这个学生不上」，
    是必须保存的信息，所以两个值都写入，不因为 False 就删行。
    要取消例外请用 clear_exception()。
    """
    student_id, weekday, block = int(student_id), int(weekday), int(block)
    note = str(note or "").strip()
    with conn:
        conn.execute(
            "INSERT INTO schedule_exceptions "
            "(student_id, weekday, block, has_class, note) VALUES (?, ?, ?, ?, ?) "
            "ON CONFLICT(student_id, weekday, block) DO UPDATE SET "
            "has_class=excluded.has_class, note=excluded.note",
            (student_id, weekday, block, 1 if has_class else 0, note),
        )
    return {"ok": True}


def clear_exception(conn: sqlite3.Connection, student_id, weekday, block) -> dict:
    """取消某人的某条例外，恢复为跟随班级课表。"""
    with conn:
        cur = conn.execute(
            "DELETE FROM schedule_exceptions "
            "WHERE student_id=? AND weekday=? AND block=?",
            (int(student_id), int(weekday), int(block)),
        )
        return {"deleted": cur.rowcount}


SCHEDULE_CSV_HEADER = ["班级", "星期", "大课", "周次", "课程"]

_WEEKDAY_WORDS = {
    "周一": 1, "星期一": 1, "礼拜一": 1, "一": 1, "1": 1,
    "周二": 2, "星期二": 2, "礼拜二": 2, "二": 2, "2": 2,
    "周三": 3, "星期三": 3, "礼拜三": 3, "三": 3, "3": 3,
    "周四": 4, "星期四": 4, "礼拜四": 4, "四": 4, "4": 4,
    "周五": 5, "星期五": 5, "礼拜五": 5, "五": 5, "5": 5,
    "周六": 6, "星期六": 6, "礼拜六": 6, "六": 6, "6": 6,
    "周日": 7, "周天": 7, "星期日": 7, "星期天": 7, "日": 7, "天": 7, "7": 7,
}


def _parse_weekday(token) -> int | None:
    return _WEEKDAY_WORDS.get(str(token or "").strip().lower().replace(" ", ""))


def _parse_block(conn, token) -> int | None:
    """大课可以写序号（1-4），也可以写名称，还能写「上①/下②」这类简写。

    简写的对应关系：上①→1 上②→2 下①→3 下②→4。
    写出的序号必须真实存在，否则视为错误（避免 99 这种脏值进库）。
    """
    raw = str(token or "").strip()
    if not raw:
        return None
    valid = {b["idx"] for b in list_blocks(conn)}
    if raw.isdigit():
        idx = int(raw)
        return idx if idx in valid else None
    circled = "①②③④⑤⑥⑦⑧⑨⑩"
    for i, ch in enumerate(circled, start=1):
        if ch in raw:
            if raw.startswith("上"):
                idx = i            # 上①→1 上②→2
            elif raw.startswith("下"):
                idx = i + 2        # 下①→3 下②→4
            else:
                idx = i
            return idx if idx in valid else None
    for b in list_blocks(conn):
        if raw == b["name"] or raw in b["name"]:
            return b["idx"]
    return None


def export_schedule_csv(conn: sqlite3.Connection) -> str:
    """导出课表为「一行一个课时」的 CSV，方便在 Excel 里改或换电脑搬。"""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(SCHEDULE_CSV_HEADER)
    rows = conn.execute(
        "SELECT cls, weekday, block, weeks, note FROM schedule "
        "WHERE has_class=1 ORDER BY cls, weekday, block"
    ).fetchall()
    for r in rows:
        writer.writerow([r["cls"], r["weekday"], r["block"],
                         r["weeks"] or "", r["note"]])
    # 个别学生例外也一并导出，格式在备注里标明，避免丢失
    for r in conn.execute(
        "SELECT e.student_id, s.name, e.weekday, e.block, e.has_class, e.note "
        "FROM schedule_exceptions e LEFT JOIN students s ON s.id = e.student_id "
        "ORDER BY e.student_id, e.weekday, e.block"
    ).fetchall():
        kind = "额外有课" if r["has_class"] else "这节不上"
        writer.writerow([f"@{r['name'] or r['student_id']}", r["weekday"], r["block"],
                         "", f"{kind}｜{r['note'] or ''}"])
    return "﻿" + buf.getvalue()


def import_schedule_csv(conn: sqlite3.Connection, text: str,
                        dry_run: bool = False) -> dict:
    """导入课表。采用「合并」而不是覆盖——避免手一滑把已录的课表清掉。

    格式：班级,星期,大课,课程；星期支持 1-7 或「周三」，大课支持序号或名称。
    班级写成 @姓名 的行会被当作「个别学生例外」。

    dry_run=True 时只解析不写库，用于在 Excel 的多个工作表里挑出课表那一张。
    """
    rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return {"added": 0, "exceptions": 0, "errors": [], "parsed": 0}

    # 跳过表头
    if rows and any(h in rows[0][0] for h in ("班级", "班", "cls")):
        rows = rows[1:]

    by_name = {r["name"]: r["id"] for r in
               conn.execute("SELECT id, name FROM students")}

    added = exceptions = parsed = 0
    errors = []
    with conn:
        for i, row in enumerate(rows, start=1):
            if len(row) < 3:
                errors.append(f"第 {i} 行字段不足：{row}")
                continue
            target = row[0].strip()
            weekday = _parse_weekday(row[1])
            block = _parse_block(conn, row[2])
            # 5 列及以上：第 4 列是周次、第 5 列是课程名
            # 正好 4 列：按旧格式，第 4 列就是课程名（周次视为每周）
            if len(row) >= 5:
                weeks_raw = row[3]
                note = row[4].strip()
            else:
                weeks_raw = ""
                note = row[3].strip() if len(row) > 3 else ""
            weeks = normalize_weeks(weeks_raw)
            if weekday is None:
                errors.append(f"第 {i} 行星期无法识别：{row[1]!r}")
                continue
            if block is None:
                errors.append(f"第 {i} 行大课无法识别：{row[2]!r}")
                continue
            parsed += 1
            if dry_run:
                continue

            if target.startswith("@"):
                name = target[1:].strip()
                sid = by_name.get(name)
                if sid is None:
                    errors.append(f"第 {i} 行找不到学生：{name}")
                    continue
                has_class = 1
                if "｜" in note or "|" in note:
                    kind, _, rest = note.replace("|", "｜").partition("｜")
                    has_class = 0 if "不上" in kind else 1
                    note = rest.strip()
                conn.execute(
                    "INSERT INTO schedule_exceptions "
                    "(student_id, weekday, block, has_class, note) VALUES (?, ?, ?, ?, ?) "
                    "ON CONFLICT(student_id, weekday, block) DO UPDATE SET "
                    "has_class=excluded.has_class, note=excluded.note",
                    (sid, weekday, block, has_class, note),
                )
                exceptions += 1
            else:
                conn.execute(
                    "INSERT INTO schedule (cls, weekday, block, has_class, weeks, note) "
                    "VALUES (?, ?, ?, 1, ?, ?) "
                    "ON CONFLICT(cls, weekday, block) DO UPDATE SET "
                    "has_class=1, weeks=excluded.weeks, note=excluded.note",
                    (target, weekday, block, weeks, note),
                )
                added += 1

    return {"parsed": parsed, "added": added, "exceptions": exceptions,
            "errors": errors[:20]}


def clear_schedule(conn: sqlite3.Connection, with_exceptions: bool = True) -> dict:
    with conn:
        cur = conn.execute("DELETE FROM schedule")
        n_exc = 0
        if with_exceptions:
            n_exc = conn.execute("DELETE FROM schedule_exceptions").rowcount
        return {"deleted": cur.rowcount, "exceptions": n_exc}


# ---------------------------------------------------------------- 教学周

def get_term_start(conn: sqlite3.Connection) -> str:
    """第 1 周周一的日期，"YYYY-MM-DD"；没设置返回空串。"""
    row = conn.execute("SELECT value FROM meta WHERE key='term_start'").fetchone()
    return (row["value"] or "") if row else ""


def set_term_start(conn: sqlite3.Connection, ymd) -> dict:
    value = str(ymd or "").strip()
    if value:
        try:
            datetime.strptime(value, "%Y-%m-%d")
        except ValueError as exc:
            raise ValueError("日期格式应为 YYYY-MM-DD") from exc
    with conn:
        conn.execute(
            "INSERT INTO meta (key, value) VALUES ('term_start', ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (value,),
        )
    return {"term_start": value}


def current_week(conn: sqlite3.Connection, now=None) -> int | None:
    """现在是第几教学周。没设学期开始日期、或今天还没开学时返回 None。"""
    start = get_term_start(conn)
    if not start:
        return None
    try:
        monday = datetime.strptime(start, "%Y-%m-%d").date()
    except ValueError:
        return None
    today = (now or datetime.now()).date()
    days = (today - monday).days
    if days < 0:
        return None
    return days // 7 + 1


def format_weeks(weeks) -> str:
    """{1,2,3,5} -> "1-3,5"；空集合 -> ""（表示每周都有课）。"""
    nums = sorted({int(w) for w in weeks if int(w) > 0})
    if not nums:
        return ""
    out, start, prev = [], nums[0], nums[0]
    for w in nums[1:]:
        if w == prev + 1:
            prev = w
            continue
        out.append(f"{start}-{prev}" if start != prev else str(start))
        start = prev = w
    out.append(f"{start}-{prev}" if start != prev else str(start))
    return ",".join(out)


def parse_week_set(text) -> set:
    """"1-8,10-16" / "1,3,5" -> 周次集合。"""
    raw = str(text or "").replace(" ", "")
    raw = raw.replace("周", "").replace("第", "")
    weeks = set()
    for part in re.split(r"[,，、;；]+", raw):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)-(\d+)", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a > b:
                a, b = b, a
            weeks.update(range(a, b + 1))
        elif part.isdigit():
            weeks.add(int(part))
    return weeks


def normalize_weeks(text) -> str:
    """把用户/教务文件里的周次写法统一成紧凑形式。

    支持："1-16"、"1-8,10-16"、"1,3,5"、"单"、"双"、空（= 每周）。
    """
    raw = str(text or "").strip().replace(" ", "")
    if not raw:
        return ""
    if raw in ("单", "单周", "单数周"):
        return "单"
    if raw in ("双", "双周", "双数周"):
        return "双"
    return format_weeks(parse_week_set(raw))


def weeks_match(weeks_text, week) -> bool:
    """这一周该时段是否有课。

    空 = 每周都有。算不出当前是第几周（week 为 None）时按「有课」处理，
    宁可少抽一个也不要抽到正在上课的人。
    """
    w = str(weeks_text or "").strip()
    if not w:
        return True
    if week is None:
        return True
    if w == "单":
        return week % 2 == 1
    if w == "双":
        return week % 2 == 0
    return week in parse_week_set(w)


def set_student_schedule(conn: sqlite3.Connection, student_id, weekday, block,
                         has_class=True, weeks="", note="") -> dict:
    """写入某学生的个人课表行（个人课表存在时，完全替代班级课表）。"""
    student_id, weekday, block = int(student_id), int(weekday), int(block)
    with conn:
        if not has_class and not str(note or "").strip():
            cur = conn.execute(
                "DELETE FROM student_schedule "
                "WHERE student_id=? AND weekday=? AND block=?",
                (student_id, weekday, block),
            )
            return {"deleted": cur.rowcount}
        conn.execute(
            "INSERT INTO student_schedule "
            "(student_id, weekday, block, has_class, weeks, note) "
            "VALUES (?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(student_id, weekday, block) DO UPDATE SET "
            "has_class=excluded.has_class, weeks=excluded.weeks, note=excluded.note",
            (student_id, weekday, block, 1 if has_class else 0,
             normalize_weeks(weeks), str(note or "").strip()),
        )
    return {"ok": True}


def clear_student_schedule(conn: sqlite3.Connection, student_id=None) -> dict:
    with conn:
        if student_id is None:
            cur = conn.execute("DELETE FROM student_schedule")
        else:
            cur = conn.execute("DELETE FROM student_schedule WHERE student_id=?",
                               (int(student_id),))
        return {"deleted": cur.rowcount}


def schedule_targets(conn: sqlite3.Connection) -> dict:
    """课表页用来展示：班级课表有哪些班、哪些学生有个人课表。"""
    classes = [dict(r) for r in conn.execute(
        "SELECT cls, COUNT(*) AS slots FROM schedule WHERE has_class=1 "
        "GROUP BY cls ORDER BY cls")]
    return {"classes": classes, "own": students_with_own_schedule(conn)}


def students_with_own_schedule(conn: sqlite3.Connection) -> list:
    rows = conn.execute(
        "SELECT DISTINCT ss.student_id, s.name, s.sid, s.cls "
        "FROM student_schedule ss LEFT JOIN students s ON s.id = ss.student_id "
        "ORDER BY s.cls, s.name"
    ).fetchall()
    return [dict(r) for r in rows]


def _student_block_map(conn: sqlite3.Connection, weekday: int,
                       week: int | None = None) -> dict:
    """指定星期几和第几教学周，每个学生「哪些大课有课」= 班级课表 叠加 个别例外。

    必须按 weekday 过滤——否则周三录的课表会在周四、周六也生效。
    还要按周次过滤：时段有时有课、有时没课的情况全靠这里。
    """
    by_class = {}
    for r in conn.execute(
        "SELECT cls, block, weeks FROM schedule WHERE has_class=1 AND weekday=?",
        (weekday,),
    ):
        if weeks_match(r["weeks"], week):
            by_class.setdefault(r["cls"], set()).add(r["block"])

    # 个人课表：只要这个学生在 student_schedule 里有记录，就完全用它，不看班级课表
    out = {}
    has_own = set()
    for r in conn.execute(
        "SELECT student_id, block, has_class, weeks FROM student_schedule "
        "WHERE weekday=?", (weekday,),
    ):
        has_own.add(r["student_id"])
        if r["has_class"] and weeks_match(r["weeks"], week):
            out.setdefault(r["student_id"], set()).add(r["block"])

    for r in conn.execute("SELECT id, cls FROM students WHERE enabled=1"):
        if r["id"] in has_own:
            out.setdefault(r["id"], set())          # 有自己的课表，班级的作废
        else:
            out[r["id"]] = set(by_class.get(r["cls"], set()))

    for r in conn.execute(
        "SELECT student_id, block, has_class FROM schedule_exceptions "
        "WHERE weekday=?", (weekday,),
    ):
        if r["student_id"] not in out:
            continue
        if r["has_class"]:
            out[r["student_id"]].add(r["block"])
        else:
            out[r["student_id"]].discard(r["block"])
    return out


def student_busy_reason(student_blocks, blocks, now, min_minutes):
    """判断某个学生此刻为什么不能抽。

    返回 None 表示空闲；否则返回 (原因, 大课, 剩余分钟)：
      ("in_class", 大课, 0)   现在就在这个有课的大课区间里
      ("soon", 大课, 分钟)    下一个「自己有课」的大课不足 min_minutes 分钟

    大课区间含内部那 10 分钟课间，所以「课间算上课期间」由区间天然覆盖。
    """
    cur = now.hour * 60 + now.minute

    for b in blocks:
        if b["idx"] in student_blocks and \
                _minute_of_day(b["start"]) <= cur < _minute_of_day(b["end"]):
            return ("in_class", b, 0)

    for b in blocks:
        start = _minute_of_day(b["start"])
        if start <= cur or b["idx"] not in student_blocks:
            continue
        left = start - cur
        if left < min_minutes:
            return ("soon", b, left)
        return None          # 第一个「自己的」未来大课已经够远，就算空闲
    return None


def pool_status(conn: sqlite3.Connection, now=None, min_minutes=None,
                ignore_schedule=False) -> dict:
    """抽取池 + 排除原因构成。界面靠这个解释「为什么人这么少」。"""
    now = now or datetime.now()
    if min_minutes is None:
        min_minutes = get_min_talk_minutes(conn)

    blocks = list_blocks(conn)
    use_schedule = schedule_configured(conn) and not ignore_schedule
    week = current_week(conn, now)
    bmap = _student_block_map(conn, now.isoweekday(), week) if use_schedule else {}

    available, in_class, soon, drawn = [], 0, 0, 0
    soon_detail = None
    reasons = {}          # 学生 id -> "in_class" / "soon"，供界面标注原因

    for r in conn.execute(
        "SELECT id, name, sid, cls, drawn FROM students WHERE enabled=1 "
        "ORDER BY cls, sid, id"
    ).fetchall():
        if r["drawn"]:
            drawn += 1
            continue
        reason = None
        if use_schedule:
            reason = student_busy_reason(
                bmap.get(r["id"], set()), blocks, now, min_minutes
            )
        if reason is None:
            available.append(dict(r))
        elif reason[0] == "in_class":
            in_class += 1
            reasons[r["id"]] = "in_class"
        else:
            soon += 1
            reasons[r["id"]] = "soon"
            if soon_detail is None:
                soon_detail = {
                    "name": r["name"],
                    "minutes": reason[2],
                    "block": reason[1]["name"],
                    "block_start": reason[1]["start"],
                }

    disabled = conn.execute(
        "SELECT COUNT(*) FROM students WHERE enabled=0"
    ).fetchone()[0]

    return {
        "available": available,
        "counts": {
            "available": len(available),
            "in_class": in_class,
            "soon": soon,
            "drawn": drawn,
            "disabled": disabled,
        },
        # 本学期还没谈过的人数（与时间无关）
        "untalked": len(available) + in_class + soon,
        "reasons": reasons,
        "schedule_active": use_schedule,
        "min_talk_minutes": min_minutes,
        "week": week,
        "term_start": get_term_start(conn),
        "own_schedule_count": len(students_with_own_schedule(conn)),
        # 课表里用了周次但没设学期开始日期 → 按每周有课保守处理，界面要提醒
        "week_unresolved": bool(
            week is None and conn.execute(
                "SELECT 1 FROM schedule WHERE has_class=1 AND weeks <> '' LIMIT 1"
            ).fetchone()
        ),
        "soon_detail": soon_detail,
        "now": now.strftime("%Y-%m-%d %H:%M:%S"),
        "weekday": now.isoweekday(),
    }


def current_block(conn: sqlite3.Connection, now=None) -> dict | None:
    """现在正处于哪个大课的区间内（用于界面显示）。"""
    now = now or datetime.now()
    cur = now.hour * 60 + now.minute
    for b in list_blocks(conn):
        if _minute_of_day(b["start"]) <= cur < _minute_of_day(b["end"]):
            return dict(b)
    return None


def next_block(conn: sqlite3.Connection, now=None) -> dict | None:
    """现在之后的下一个大课（用于显示「下节课 HH:MM 开始」）。

    与 next_free_hint 不同：那个返回的是"什么时候下课/有人空闲"，
    这个返回的是"下一节课什么时候开始"，别混用。
    """
    now = now or datetime.now()
    cur = now.hour * 60 + now.minute
    for b in list_blocks(conn):
        if _minute_of_day(b["start"]) > cur:
            return dict(b)
    return None


def next_free_hint(conn: sqlite3.Connection, now=None) -> dict | None:
    """给「此刻无人可抽」时用：提示下一段空闲大约从什么时候开始。"""
    now = now or datetime.now()
    blocks = list_blocks(conn)
    if not blocks:
        return None
    cur = now.hour * 60 + now.minute
    for b in blocks:
        if _minute_of_day(b["end"]) > cur:
            return {"at": b["end"], "after": b["name"]}
    return None


# ---------------------------------------------------------------- 名单维护

def add_student(
    conn: sqlite3.Connection, name: str, sid: str = "", cls: str = ""
) -> dict:
    name = (name or "").strip()
    if not name:
        raise ValueError("姓名不能为空")
    with conn:
        cur = conn.execute(
            "INSERT INTO students (name, sid, cls, enabled, drawn, created_at) "
            "VALUES (?, ?, ?, 1, 0, ?)",
            (name, (sid or "").strip(), (cls or "").strip(), now_str()),
        )
        return {"id": cur.lastrowid, "name": name}


def update_student(
    conn: sqlite3.Connection,
    student_id: int,
    name: str | None = None,
    sid: str | None = None,
    cls: str | None = None,
) -> dict:
    fields, values = [], []
    for col, val in (("name", name), ("sid", sid), ("cls", cls)):
        if val is not None:
            fields.append(f"{col}=?")
            values.append(val.strip())
    if not fields:
        return {"updated": 0}
    if name is not None and not name.strip():
        raise ValueError("姓名不能为空")
    with conn:
        values.append(student_id)
        cur = conn.execute(
            f"UPDATE students SET {', '.join(fields)} WHERE id=?", values
        )
        return {"updated": cur.rowcount}


def set_enabled(conn: sqlite3.Connection, student_id: int, enabled: bool) -> dict:
    """手动启用/停用。停用只是不参与抽取，不是删除。"""
    with conn:
        cur = conn.execute(
            "UPDATE students SET enabled=? WHERE id=?",
            (1 if enabled else 0, student_id),
        )
        return {"updated": cur.rowcount}


def set_enabled_many(conn: sqlite3.Connection, ids, enabled: bool) -> dict:
    """批量启用/停用，一个事务搞定。"""
    ids = [int(i) for i in (ids or [])]
    if not ids:
        return {"updated": 0}
    marks = ",".join("?" * len(ids))
    with conn:
        cur = conn.execute(
            f"UPDATE students SET enabled=? WHERE id IN ({marks})",
            [1 if enabled else 0] + ids,
        )
        return {"updated": cur.rowcount}


def delete_students(conn: sqlite3.Connection, ids) -> dict:
    """批量永久删除。抽取历史保留姓名快照，只断开学生引用。"""
    ids = [int(i) for i in (ids or [])]
    if not ids:
        return {"deleted": 0}
    marks = ",".join("?" * len(ids))
    with conn:
        conn.execute(
            f"UPDATE draws SET student_id=NULL WHERE student_id IN ({marks})", ids
        )
        cur = conn.execute(f"DELETE FROM students WHERE id IN ({marks})", ids)
        return {"deleted": cur.rowcount}


def delete_student(conn: sqlite3.Connection, student_id: int) -> dict:
    """永久删除该学生。

    抽取历史里冗余存了姓名快照，所以历史记录不会因此变成空白。
    """
    with conn:
        row = conn.execute(
            "SELECT name FROM students WHERE id=?", (student_id,)
        ).fetchone()
        if row is None:
            raise ValueError("该学生已不存在")
        conn.execute("DELETE FROM students WHERE id=?", (student_id,))
        # 已抽出的历史记录保留，但断开学生引用，撤回按钮将失效
        conn.execute(
            "UPDATE draws SET student_id=NULL WHERE student_id=?", (student_id,)
        )
        return {"name": row["name"]}


def clear_students(conn: sqlite3.Connection) -> dict:
    with conn:
        cur = conn.execute("DELETE FROM students")
        conn.execute("UPDATE draws SET student_id=NULL")
        return {"deleted": cur.rowcount}


# ---------------------------------------------------------------- CSV 导入

def decode_bytes(raw: bytes) -> str:
    """Excel 另存的 CSV 在中文 Windows 上常是 GBK，优先按 UTF-8 试，再退 GBK。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _sniff_delimiter(line: str) -> str:
    counts = {d: line.count(d) for d in (",", "\t", ";")}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def _map_header(cells: list[str]) -> dict | None:
    """若首行是表头，返回列名到下标的映射，否则返回 None。"""
    mapping = {}
    for idx, cell in enumerate(cells):
        token = cell.strip().lower().replace(" ", "")
        for field, aliases in HEADER_ALIASES.items():
            if token in {a.lower() for a in aliases}:
                mapping.setdefault(field, idx)
    if "name" in mapping:
        return mapping
    return None


def parse_students_text(text: str) -> list[dict]:
    """解析名单文本，自动处理分隔符、表头、有无学号等情况。"""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return []

    delimiter = _sniff_delimiter(lines[0])
    rows = list(csv.reader(lines, delimiter=delimiter))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []

    mapping = _map_header(rows[0])
    if mapping:
        body = rows[1:]
    else:
        body = rows
        # 无表头时嗅探列序：首列整列都是长数字则认为是「学号,姓名,班级」
        first_col = [r[0].strip() for r in body if r and r[0].strip()]
        looks_like_id = first_col and all(
            c.isdigit() and len(c) >= 6 for c in first_col
        )
        if looks_like_id:
            mapping = {"sid": 0, "name": 1, "cls": 2}
        else:
            mapping = {"name": 0, "sid": 1, "cls": 2}

    def cell(row: list[str], field: str) -> str:
        idx = mapping.get(field)
        if idx is None or idx >= len(row):
            return ""
        return row[idx].strip().strip('"').strip()

    out = []
    for row in body:
        if not row:
            continue
        if len(row) == 1:
            # 单列纯名单，兼容「1. 张三」这类行首序号
            name = _NUMBERING.sub("", row[0].strip()).strip()
            sid = cls = ""
        else:
            name = cell(row, "name")
            sid = cell(row, "sid")
            cls = cell(row, "cls")
        if not name:
            continue
        out.append({"name": name, "sid": sid, "cls": cls})
    return out


def _cell_to_str(value) -> str:
    """把 Excel 单元格值规整成文本。

    学号常被 Excel 存成数值，2024240101 会变成 2024240101.0，
    这里统一去掉多余的小数尾巴。
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_to_text(rows: list[list]) -> str:
    """把二维表转成 CSV 文本，同时跳过表头之前的标题行。

    学校发的名单常在首行放「XX班学生名单」这类标题，
    真正表头在第 2 行，这里自动定位。
    """
    cleaned = [[_cell_to_str(c) for c in row] for row in rows]
    cleaned = [r for r in cleaned if any(c for c in r)]
    if not cleaned:
        return ""

    start = 0
    for idx, row in enumerate(cleaned[:10]):
        if _map_header(row) is not None:
            start = idx
            break
    else:
        # 没有可识别表头时，跳过明显只有一格的标题行
        for idx, row in enumerate(cleaned[:10]):
            if sum(1 for c in row if c) >= 2:
                start = idx
                break

    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerows(cleaned[start:])
    return buf.getvalue()


def read_table_file(path: str | Path, probe=None) -> tuple[str, str]:
    """读取 .csv / .xlsx，返回 (文本, 说明)。

    .xlsx 会扫描所有工作表，选「最像目标数据」的那一张，这样即使第一张是封面页
    也能正确取到内容。判断依据由 probe 决定：
      - 不传 probe：按解析出的学生数（导入名单用）
      - 导入课表时传课表专用的 probe，否则会按学生数挑到错误的表
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in (".csv", ".txt"):
        text = decode_bytes(path.read_bytes())
        return text, path.name

    if suffix == ".xls":
        import jiaowu

        try:
            rows, _ = jiaowu.read_rows(path)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"这个 .xls 文件读不了：{exc}") from exc
        return _rows_to_text(rows), path.name

    if suffix != ".xlsx":
        raise ValueError(
            f"不支持的文件类型：{suffix}。请提供 .xlsx / .xls / .csv 文件。")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("缺少 openpyxl，无法读取 Excel 文件") from exc

    score = probe or (lambda text: len(parse_students_text(text)))
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        best_text, best_count, best_name = "", -1, ""
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            text = _rows_to_text(rows)
            count = score(text)
            if count > best_count:
                best_text, best_count, best_name = text, count, ws.title
        return best_text, f"{path.name} / 工作表「{best_name}」"
    finally:
        wb.close()


def _extract_sid_from_filename(path: Path) -> str:
    """从「学生个人课表_24201320.xls」里取出学号。"""
    stem = path.stem
    for part in reversed(stem.replace("-", "_").split("_")):
        if part.isdigit() and len(part) >= 4:
            return part
    return ""


def import_jiaowu_grid(conn: sqlite3.Connection, grid: dict,
                       target: str = "class", cls: str | None = None,
                       path: Path | None = None) -> dict:
    """把教务二维课表写进库。

    target:
      "class"   写成班级课表（默认）——文件里的「班级：」或参数 cls
      "student" 写成学生个人课表（按学号匹配，对不上就报错）

    默认按班级：教务导出的虽是「学生个人课表」，但辅导员通常是要给整个班用。
    曾经默认「能匹配到学生就写成个人课表」，结果只排除了那一个人，
    同一班其他同学照旧能被抽到，很隐蔽。
    """
    meta = grid.get("meta") or {}
    cls_name = (cls or meta.get("班级") or "").strip()
    if not cls_name and path is not None:
        # 兜底：班级也可能只出现在文件名里，如「课表打印-242013班.xls」
        m = re.search(r"(?<!\d)(\d{4,8})\s*班", path.stem)
        if m:
            cls_name = m.group(1)
    sid = (meta.get("学号") or (path and _extract_sid_from_filename(path)) or "").strip()

    student_id = None
    if sid:
        row = conn.execute("SELECT id, name FROM students WHERE sid=? LIMIT 1",
                           (sid,)).fetchone()
        if row:
            student_id = row["id"]

    if target == "student":
        if student_id is None:
            return {"error": f"按学号 {sid!r} 在名单里找不到学生，无法写成个人课表。",
                    "added": 0, "exceptions": 0, "errors": []}
    else:
        if not cls_name:
            return {"error": "文件里没有「班级：」信息，无法确定写到哪个班。",
                    "added": 0, "exceptions": 0, "errors": []}
        # 关键的防护：班级在名单里不存在的话，导入进去也完全不生效，
        # 必须报错而不是默默成功
        n = conn.execute("SELECT COUNT(*) FROM students WHERE cls=?",
                         (cls_name,)).fetchone()[0]
        if n == 0:
            names = [r["cls"] for r in conn.execute(
                "SELECT DISTINCT cls FROM students WHERE cls <> '' ORDER BY cls")]
            return {"error": f"课表里的班级「{cls_name}」在名单里不存在，"
                             f"导进去不会生效。名单里的班级是：{'、'.join(names) or '（名单为空）'}。"
                             f"如果确认要用，请先在名单里把班级名改成一致。",
                    "added": 0, "exceptions": 0, "errors": []}
        student_id = None

    db.ensure_blocks(conn, max(4, int(grid.get("max_slot") or 4)))

    added = 0
    with conn:
        if student_id is not None:
            conn.execute("DELETE FROM student_schedule WHERE student_id=?", (student_id,))
            for (wd, blk), cell in grid["cells"].items():
                conn.execute(
                    "INSERT INTO student_schedule "
                    "(student_id, weekday, block, has_class, weeks, note) "
                    "VALUES (?, ?, ?, 1, ?, ?)",
                    (student_id, wd, blk, cell["weeks"], cell["note"]),
                )
                added += 1
        else:
            for (wd, blk), cell in grid["cells"].items():
                conn.execute(
                    "INSERT INTO schedule (cls, weekday, block, has_class, weeks, note) "
                    "VALUES (?, ?, ?, 1, ?, ?) "
                    "ON CONFLICT(cls, weekday, block) DO UPDATE SET "
                    "has_class=1, weeks=excluded.weeks, note=excluded.note",
                    (cls_name, wd, blk, cell["weeks"], cell["note"]),
                )
                added += 1

    who = ""
    if student_id is not None:
        row = conn.execute("SELECT name FROM students WHERE id=?", (student_id,)).fetchone()
        who = f"学生 {row['name']}（个人课表，将覆盖班级课表）"
    else:
        who = f"班级 {cls_name}"
    return {"parsed": added, "added": added, "exceptions": 0, "errors": [],
            "source": getattr(path, "name", "") or "教务课表",
            "target": who, "max_slot": grid.get("max_slot")}


def import_schedule_from_file(conn: sqlite3.Connection, path: str | Path,
                              target: str = "class", cls: str | None = None) -> dict:
    """导入课表：自动识别教务导出的二维课表，也支持「一行一个课时」的 CSV。

    二维课表按 target 决定写到班级还是学生；CSV 沿用原来的合并逻辑。
    """
    import jiaowu

    path = Path(path)
    if path.suffix.lower() in (".xls", ".xlsx"):
        try:
            if jiaowu.looks_like_jiaowu(jiaowu.read_rows(path)[0]):
                grid = jiaowu.parse_file(path)
                return import_jiaowu_grid(conn, grid, target=target, cls=cls, path=path)
        except ValueError:
            raise
        except Exception:
            pass          # 不是这种格式就退回下面的通用读取
    return _import_tidy_schedule(conn, path)


def _import_tidy_schedule(conn: sqlite3.Connection, path: Path) -> dict:
    """原来的「一行一个课时」CSV/xlsx 导入路径。"""
    text, source = read_table_file(
        path, probe=lambda t: import_schedule_csv(conn, t, dry_run=True)["parsed"]
    )
    result = import_schedule_csv(conn, text)
    result["source"] = source
    return result


def import_students_from_file(conn: sqlite3.Connection, path: str | Path) -> dict:
    text, source = read_table_file(path)
    result = import_students(conn, text)
    result["source"] = source
    return result


def import_students(conn: sqlite3.Connection, text: str) -> dict:
    """导入名单，跳过与现有记录完全重复的条目。"""
    parsed = parse_students_text(text)
    existing = {
        (r["name"], r["sid"], r["cls"])
        for r in conn.execute("SELECT name, sid, cls FROM students")
    }
    seen_sid: dict[str, str] = {}
    seen_name: dict[str, str] = {}

    added, skipped, warnings = 0, 0, []
    with conn:
        for item in parsed:
            key = (item["name"], item["sid"], item["cls"])
            if key in existing:
                skipped += 1
                continue
            conn.execute(
                "INSERT INTO students (name, sid, cls, enabled, drawn, created_at) "
                "VALUES (?, ?, ?, 1, 0, ?)",
                (item["name"], item["sid"], item["cls"], now_str()),
            )
            existing.add(key)
            added += 1

            if item["sid"]:
                if item["sid"] in seen_sid and seen_sid[item["sid"]] != item["name"]:
                    warnings.append(
                        f"学号 {item['sid']} 对应多个姓名："
                        f"{seen_sid[item['sid']]} / {item['name']}"
                    )
                seen_sid[item["sid"]] = item["name"]
            if item["name"] in seen_name and item["sid"] and seen_name[item["name"]] != item["sid"]:
                warnings.append(
                    f"姓名 {item['name']} 出现多个学号："
                    f"{seen_name[item['name']]} / {item['sid']}"
                )
            seen_name[item["name"]] = item["sid"]

    return {
        "parsed": len(parsed),
        "added": added,
        "skipped": skipped,
        "warnings": warnings[:20],
    }


# ---------------------------------------------------------------- 导出

def export_csv_text(conn: sqlite3.Connection, round: int | None = None) -> str:
    """导出抽取记录为 CSV。

    日期与时间分成两列，便于按天统计。
    带 UTF-8 BOM，Excel 双击打开不会乱码。
    """
    rows = history(conn, round)
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(CSV_HEADER)
    for r in rows:
        day, clock = split_stamp(r["drawn_at"])
        writer.writerow(
            [r["seq"], day, clock, r["round"], r["name"], r["sid"], r["cls"]]
        )
    return "\ufeff" + buf.getvalue()


def export_backup_text(conn: sqlite3.Connection) -> str:
    payload = {
        "app": "随机抽人",
        "version": 1,
        "exported_at": now_str(),
        "round": get_round(conn),
        "students": list_students(conn),
        "draws": [
            {k: v for k, v in r.items() if k != "id"} for r in history(conn, None)
        ],
        # 课表数据一并备份，换电脑时不丢
        "blocks": list_blocks(conn),
        "schedule": [
            dict(r) for r in conn.execute(
                "SELECT cls, weekday, block, has_class, weeks, note FROM schedule")
        ],
        "schedule_exceptions": [
            dict(r) for r in conn.execute(
                "SELECT student_id, weekday, block, has_class, note "
                "FROM schedule_exceptions")
        ],
        "min_talk_minutes": get_min_talk_minutes(conn),
        "term_start": get_term_start(conn),      # 漏了它，按周判断会退化成"每周有课"
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def import_backup_text(conn: sqlite3.Connection, text: str) -> dict:
    """从备份恢复（整体覆盖当前数据）。"""
    data = json.loads(text)
    students = data.get("students") or []
    draws = data.get("draws") or []
    if not isinstance(students, list):
        raise ValueError("备份文件格式不正确")

    with conn:
        conn.execute("DELETE FROM students")
        conn.execute("DELETE FROM draws")
        for s in students:
            conn.execute(
                "INSERT INTO students (name, sid, cls, enabled, drawn, drawn_at, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    s.get("name", ""),
                    s.get("sid", ""),
                    s.get("cls", ""),
                    1 if s.get("enabled", 1) else 0,
                    1 if s.get("drawn", 0) else 0,
                    s.get("drawn_at"),
                    s.get("created_at") or now_str(),
                ),
            )
        for d in draws:
            conn.execute(
                "INSERT INTO draws (student_id, seq, round, drawn_at, name, sid, cls) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    d.get("student_id"),
                    d.get("seq", 0),
                    d.get("round", 1),
                    d.get("drawn_at", ""),
                    d.get("name", ""),
                    d.get("sid", ""),
                    d.get("cls", ""),
                ),
            )
        set_round(conn, int(data.get("round", 1)))

        if data.get("blocks"):
            conn.execute("DELETE FROM blocks")
            for b in data["blocks"]:
                conn.execute(
                    "INSERT INTO blocks (idx, name, start, end) VALUES (?, ?, ?, ?)",
                    (b.get("idx"), b.get("name", ""), b.get("start", "08:00"),
                     b.get("end", "08:45")),
                )
        if data.get("schedule") is not None:
            conn.execute("DELETE FROM schedule")
            for r in data["schedule"]:
                conn.execute(
                    "INSERT INTO schedule "
                    "(cls, weekday, block, has_class, weeks, note) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (r.get("cls", ""), r.get("weekday", 1), r.get("block", 1),
                     1 if r.get("has_class", 1) else 0, r.get("weeks", ""),
                     r.get("note", "")),
                )
        if data.get("schedule_exceptions") is not None:
            conn.execute("DELETE FROM schedule_exceptions")
            for r in data["schedule_exceptions"]:
                conn.execute(
                    "INSERT INTO schedule_exceptions "
                    "(student_id, weekday, block, has_class, note) "
                    "VALUES (?, ?, ?, ?, ?)",
                    (r.get("student_id"), r.get("weekday", 1), r.get("block", 1),
                     1 if r.get("has_class", 1) else 0, r.get("note", "")),
                )
        if data.get("min_talk_minutes") is not None:
            set_min_talk_minutes(conn, data["min_talk_minutes"])
        if data.get("term_start") is not None:
            set_term_start(conn, data["term_start"])

    return {"students": len(students), "draws": len(draws)}
